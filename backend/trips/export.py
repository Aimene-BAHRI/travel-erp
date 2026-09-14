"""Read-only XLSX export using the same readiness results displayed by the API."""
from io import BytesIO
from decimal import Decimal
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone
from bookings.models import BookingTraveler
from bookings.readiness import ACTIVE, current_passport, loaded_links, trip_readiness
from finance.services import booking_financial_summary


def build_trip_workbook(trip, actor):
    report = trip_readiness(trip)
    links = list(loaded_links(BookingTraveler.objects.filter(booking__trip=trip, booking__status__in=ACTIVE)))
    details = {row["id"]: row for row in report["travelers_detail"]}
    stops = list(trip.stops.select_related("city").order_by("order"))
    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = {}
    def sheet(name, headings):
        ws = workbook.create_sheet(name)
        ws.sheet_view.rightToLeft = True
        ws.append(headings)
        sheets[name] = ws
        return ws
    main = sheet("المسافرون", ["رقم", "الحجز", "صاحب الحجز", "الاسم واللقب", "هاتف المسافر", "هاتف صاحب الحجز", "ملاحظة", "التصنيف", "السرير", "مقعد", "الجنس", "نوع الخدمة"] + [f"{stop.city.name} {stop.order} ? {label}" for stop in stops for label in ["الفندق", "الغرفة", "نوع الغرفة"]] + ["رقم الجواز", "انتهاء الجواز", "العملة", "السعر الأساسي", "التخفيض", "السعر النهائي", "المدفوع المخصص", "الباقي من السعر الفردي", "حالة الملف"])
    issues_sheet = sheet("البيانات الناقصة", ["الحجز", "المسافر", "الحقل", "الحالة", "التفاصيل"])
    hotels = sheet("توزيع الفنادق", ["المدينة", "الفندق", "الحجز", "المسافر", "التصنيف", "يستهلك سريراً"])
    rooms = sheet("توزيع الغرف", ["المدينة", "الفندق", "الغرفة", "نوع الغرفة", "السعة", "المسافر", "رقم السرير", "يستهلك سريراً"])
    for index, link in enumerate(links, 1):
        p = current_passport(link)
        row = details[str(link.pk)]
        charges = list(link.charges.all())
        base = sum((c.base_amount for c in charges), Decimal("0"))
        discount = sum((c.discount_amount for c in charges), Decimal("0"))
        final = base - discount
        paid = sum((a.amount * a.transaction.direction for a in link.payment_allocations.all()), Decimal("0"))
        assignments = list(link.trip_room_assignments.all())
        itinerary = []
        for stop in stops:
            selected = next((a for a in assignments if a.stop_id == stop.pk), None)
            source = next((c.source_price for c in charges if c.source_price_id and c.source_price.trip_hotel.stop_id == stop.pk), None)
            hotel = selected.room.trip_hotel.hotel.name if selected else source.trip_hotel.hotel.name if source else None
            itinerary += [hotel, selected.room.label if selected else None, selected.room.room_type.name if selected else source.room_type.name if source else None]
            if hotel:
                hotels.append([stop.city.name, hotel, link.booking.code, str(link.traveler), row["category"], row["bed_consuming"]])
            if selected:
                rooms.append([stop.city.name, hotel, selected.room.label, selected.room.room_type.name, selected.room.capacity, str(link.traveler), selected.bed_number, row["bed_consuming"]])
        main.append([index, link.booking.code, str(link.booking.customer), str(link.traveler), link.traveler.phone or None, link.booking.customer.phone or None, link.notes or link.traveler.notes or None, row["category"], "مع سرير" if row["bed_consuming"] else "بدون سرير", row["seat_consuming"], link.traveler.gender, link.service_type] + itinerary + [p.passport_number if p else None, p.expiry_date if p else None, link.booking.currency_id, base if charges else None, discount if charges else None, final if charges else None, paid, final - paid if charges else None, row["status"]])
        for key, state in row["fields"].items():
            if state in {"MISSING", "INVALID"}:
                issues_sheet.append([link.booking.code, str(link.traveler), key, state, " · ".join(row["issues"])])
    summary = sheet("ملخص الرحلة", ["البيان", "القيمة"])
    t, seat, quality = report["travelers"], report["seats"], report["data_quality"]
    for row in [("الرحلة", trip.name), ("المغادرة", trip.departure_date), ("العودة", trip.return_date), ("السعة", seat["capacity"]), ("المقاعد المشغولة", seat["occupied"]), ("المقاعد المتبقية", seat["remaining"]), ("تجاوز السعة", seat["overbooked_places"]), ("عدد المسافرين", t["total"]), ("ADT", t["adt"]), ("CHD", t["chd"]), ("INF", t["inf"]), ("مسافرون بسرير", t["bed_consuming"]), ("مسافرون بدون سرير", t["total"] - t["bed_consuming"]), ("ملفات مكتملة", quality["valid"]), ("ملفات غير مكتملة", quality["incomplete"]), ("جوازات مفقودة", report["missing"]["passport"]), ("جوازات غير صالحة", report["invalid"]["passport"]), ("جوازات صالحة", t["total"] - report["missing"]["passport"] - report["invalid"]["passport"]), ("تاريخ الإنشاء", timezone.now().isoformat()), ("أنشأه", str(actor)), ("مصدر البيانات", "Wakaltcom — real confirmed/completed booking travelers only"), ("حقول غير متتبعة", "Visa, ticket status, meals, manasik: no structured source recorded"), ("ملاحظة الدفعات", "Traveler paid amounts include explicit allocations only; unallocated receipts remain on the bookings sheet.")]:
        summary.append(row)
    bookings = sheet("الحجوزات", ["الحجز", "صاحب الحجز", "الفرع", "العدد المطلوب", "المسافرون الفعليون", "العملة", "الإجمالي", "المدفوع", "الباقي", "المدفوع غير المخصص", "ملاحظات"])
    for booking in trip.bookings.filter(status__in=ACTIVE).select_related("customer", "branch").prefetch_related("financial_transactions__reverses", "travelers"):
        financial = booking_financial_summary(booking)
        assigned_paid = sum((a.amount * a.transaction.direction for link in links if link.booking_id == booking.pk for a in link.payment_allocations.all()), Decimal("0"))
        bookings.append([booking.code, str(booking.customer), booking.branch.name, booking.expected_travelers, len(booking.travelers.all()), booking.currency_id, booking.total_amount, financial["paid_amount"], financial["remaining_amount"], financial["paid_amount"] - assigned_paid, booking.notes])
    workbook._sheets = [main, summary, issues_sheet, hotels, rooms, bookings]
    line = Side(style="hair", color="DCE4E8")
    for ws in workbook:
        ws.freeze_panes = "E2" if ws == main else "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if ws == main else ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        ws.row_dimensions[1].height = 34
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="155E75")
            cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    # Untrusted names/notes must never become Excel formulas.
                    cell.data_type = "s"
                cell.font = Font(name="Arial", size=10)
                cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
                cell.border = Border(bottom=line)
                if cell.row % 2 == 0: cell.fill = PatternFill("solid", fgColor="F3F7F8")
                if cell.is_date: cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, (Decimal, float)): cell.number_format = '#,##0.00;[Red](#,##0.00)'
                if cell.value in {"MISSING", "INVALID"}: cell.fill = PatternFill("solid", fgColor="FFF1CC" if cell.value == "MISSING" else "FEE2E2")
        for column in ws.columns:
            width = min(36, max(14, max(len(str(c.value or "")) for c in column) + 2))
            ws.column_dimensions[get_column_letter(column[0].column)].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
