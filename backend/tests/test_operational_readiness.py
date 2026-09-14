from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace

import pytest
from django.core.exceptions import ValidationError
from openpyxl import load_workbook

from tests.test_bookings import booking_data, create_booking, client_for
from bookings.models import BookingTraveler
from bookings.readiness import consumes_bed, consumes_seat, required_stops, trip_readiness, require_room_ready
from pricing.services import resolve_traveler_price
from reference.models import PassengerCategory
from trips.export import build_trip_workbook


@pytest.mark.parametrize('category,bed,seat,expected_seat,expected_bed', [
    ('ADT', 'REQUIRED', None, True, True),
    ('CHD', 'REQUIRED', None, True, True),
    ('CHD', 'NOT_REQUIRED', None, True, False),
    ('INF', 'NOT_REQUIRED', None, False, False),
    ('INF', 'REQUIRED', True, True, False),
    ('ADT', 'REQUIRED', False, False, True),
])
def test_seats_and_beds_are_independent(category, bed, seat, expected_seat, expected_bed):
    link = SimpleNamespace(passenger_category=SimpleNamespace(code=category), bed_requirement=bed, seat_required=seat)
    assert consumes_seat(link) is expected_seat
    assert consumes_bed(link) is expected_bed


@pytest.mark.parametrize('service,expected', [('PACK', 1), ('VISA_ONLY', 0), ('TICKET_ONLY', 0), ('CUSTOM', 0)])
def test_service_specific_stays(booking_data, service, expected):
    d = booking_data
    booking = create_booking(d)
    link = BookingTraveler.objects.create(booking=booking, traveler=d['traveler'], passenger_category=d['adult'], service_type=service, bed_requirement='REQUIRED')
    assert len(required_stops(link, [d['stop']])) == expected
    if service == 'CUSTOM':
        link.included_stops.add(d['stop'])
        assert required_stops(link, [d['stop']]) == [d['stop']]


@pytest.mark.parametrize('category,bed,expected', [('ADT', 'REQUIRED', '120000'), ('CHD', 'REQUIRED', '100000'), ('CHD', 'NOT_REQUIRED', '70000'), ('INF', 'NOT_REQUIRED', '25000')])
def test_package_prices_and_historical_price_preservation(booking_data, category, bed, expected):
    d = booking_data
    trip = d['trip']
    trip.pricing_currency = d['dzd']
    trip.child_with_bed_discount = Decimal('20000')
    trip.child_without_bed_price = Decimal('70000')
    trip.infant_price = Decimal('25000')
    price = d['hotel_price']
    price.scope = 'TRIP_PACKAGE'
    price.save(update_fields=['scope'])
    result = resolve_traveler_price(trip=trip, trip_hotel=d['trip_hotel'], room_type=d['room_type'], passenger_category=SimpleNamespace(code=category), bed_requirement=bed)
    assert result.amount == Decimal(expected)
    price.refresh_from_db()
    assert price.amount == Decimal('120000')


def test_excessive_discount_rejected(booking_data):
    d = booking_data
    d['trip'].child_with_bed_discount = Decimal('120001')
    with pytest.raises(ValidationError, match='discount'):
        resolve_traveler_price(trip=d['trip'], trip_hotel=d['trip_hotel'], room_type=d['room_type'], passenger_category=SimpleNamespace(code='CHD'), bed_requirement='REQUIRED', scope='HOTEL_STAY')


def test_names_only_export_matches_readiness_and_preserves_data(booking_data):
    d = booking_data
    booking = create_booking(d)
    d['trip'].max_travelers = 1
    d['trip'].save(update_fields=['max_travelers'])
    link = BookingTraveler.objects.create(booking=booking, traveler=d['traveler'], passenger_category=d['adult'])
    with pytest.raises(ValidationError):
        require_room_ready(link)
    before = (booking.total_amount, BookingTraveler.objects.count())
    report = trip_readiness(d['trip'])
    assert report['travelers']['total'] == 1
    assert report['missing']['passport'] == 1
    assert report['seats']['remaining'] == 0
    assert report['travelers_detail'][0]['eligible_for_room'] is False
    workbook = load_workbook(BytesIO(build_trip_workbook(d['trip'], d['owner'])))
    assert len(workbook.worksheets) == 6
    assert workbook.worksheets[0].max_row == report['travelers']['total'] + 1
    assert all(ws.sheet_view.rightToLeft for ws in workbook)
    assert workbook.worksheets[0].title == 'المسافرون'
    summary = dict(workbook.worksheets[1].values)
    assert summary['عدد المسافرين'] == report['travelers']['total']
    assert summary['المقاعد المشغولة'] == report['seats']['occupied']
    booking.refresh_from_db()
    assert before == (booking.total_amount, BookingTraveler.objects.count())
    client = client_for(d['owner'])
    response = client.get(f'/api/v1/trips/{d["trip"].pk}/readiness/')
    assert response.status_code == 200
    assert response.data['travelers'] == report['travelers']


def test_infant_does_not_use_last_seat(booking_data):
    d = booking_data
    booking = create_booking(d)
    d['trip'].max_travelers = 0
    d['trip'].save(update_fields=['max_travelers'])
    BookingTraveler.objects.create(booking=booking, traveler=d['traveler'], passenger_category=PassengerCategory.objects.get(code='inf', is_system=True), bed_requirement='NOT_REQUIRED')
    report = trip_readiness(d['trip'])
    assert report['seats']['occupied'] == 0
    assert report['seats']['overbooked_places'] == 0
    assert report['missing']['room'] == 0
